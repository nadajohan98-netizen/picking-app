"""
Lectura de un archivo de lotes subido por el administrador (Fase 5).

Acepta .xlsx y .csv. Busca una columna llamada "lote" / "codigo" / "código"
(sin importar mayúsculas); si no la encuentra, usa la primera columna.
Devuelve la lista de códigos encontrados, ya limpios y sin repetir.
"""

from __future__ import annotations

import csv
import io

from openpyxl import load_workbook

_NOMBRES_COLUMNA = {"lote", "codigo", "código", "code"}


def _elegir_columna(encabezados: list[str]) -> int:
    for i, nombre in enumerate(encabezados):
        if str(nombre).strip().lower() in _NOMBRES_COLUMNA:
            return i
    return 0


def extraer_codigos(nombre_archivo: str, contenido: bytes) -> list[str]:
    nombre = nombre_archivo.lower()
    if nombre.endswith(".csv"):
        filas = _leer_csv(contenido)
    elif nombre.endswith(".xlsx"):
        filas = _leer_xlsx(contenido)
    else:
        raise ValueError("Formato no soportado. Sube un archivo .xlsx o .csv.")

    if not filas:
        return []

    encabezados = [str(c) if c is not None else "" for c in filas[0]]
    col = _elegir_columna(encabezados)

    vistos: set[str] = set()
    codigos: list[str] = []
    for fila in filas[1:]:
        if col >= len(fila):
            continue
        valor = fila[col]
        if valor is None:
            continue
        codigo = str(valor).strip()
        if codigo and codigo not in vistos:
            vistos.add(codigo)
            codigos.append(codigo)
    return codigos


def _leer_csv(contenido: bytes) -> list[list]:
    texto = contenido.decode("utf-8-sig", errors="replace")
    return [fila for fila in csv.reader(io.StringIO(texto))]


def _leer_xlsx(contenido: bytes) -> list[list]:
    libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    hoja = libro.active
    return [list(fila) for fila in hoja.iter_rows(values_only=True)]
